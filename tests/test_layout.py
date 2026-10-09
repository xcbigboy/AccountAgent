import copy
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

from accounting_report.ai import request_json, DEEPSEEK_ENDPOINT, DEEPSEEK_MODEL
from accounting_report.config import load_config
from accounting_report.errors import ReportError
from accounting_report.layout import scan_layout, discover, make_candidate, approve
from accounting_report.pipeline import run

PROJECT = Path(__file__).resolve().parents[1]


class LayoutTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.excel = self.root / "shifted.xlsx"
        self.config_path = PROJECT / "examples/mapping.json"
        self.config = load_config(self.config_path)
        source = load_workbook(PROJECT / "examples/balance_sheet.xlsx", data_only=True)
        self.addCleanup(source.close)
        wb = Workbook()
        ws = wb.active
        ws.title = "新布局"
        original = source["资产负债表"]
        # Deliberately move every row and column; a fixed old config cannot run.
        for row in original:
            for cell in row:
                if cell.value is not None:
                    ws.cell(cell.row + 3, cell.column + 2, cell.value)
        wb.save(self.excel)
        wb.close()
        self.answer = {"sheet": "新布局", "metadata": [{"original_cell": m["cell"], "cell": self.shift(m["cell"])}
                        for m in self.config["metadata"]], "matches": [], "unresolved": []}
        for field in self.config["fields"]:
            src = field["source"]
            b = next(b for b in self.config["source"]["blocks"] if b["id"] == src["block"])
            row = next(r for r in range(b["start_row"], b["end_row"] + 1)
                       if original[f"{b['label_column']}{r}"].value in src["labels"])
            column = b["period_columns"][src["period"]]
            self.answer["matches"].append({"field": field["id"], "label_cell": self.shift(f"{b['label_column']}{row}"),
                "value_cell": self.shift(f"{column}{row}"), "header_cell": self.shift(f"{column}4"), "reason": "名称及日期对应"})
        self.inventory = scan_layout(self.excel)

    def shift(self, cell):
        col, row = coordinate_from_string(cell)
        return f"{get_column_letter(column_index_from_string(col) + 2)}{row + 3}"

    def write_proposal(self):
        with patch("accounting_report.layout.request_json", return_value=self.answer):
            proposal = discover(self.excel, self.config_path)
        path = self.root / "proposal.json"
        path.write_text(json.dumps(proposal, ensure_ascii=False), encoding="utf-8")
        return path

    def reject(self, edit):
        answer = copy.deepcopy(self.answer)
        edit(answer)
        with self.assertRaises(ReportError):
            make_candidate(self.config, self.inventory, answer)

    def test_scan_no_network_and_numeric_values_hidden(self):
        with patch("accounting_report.ai.build_opener", side_effect=AssertionError("network")):
            result = scan_layout(self.excel)
        cells = {c["cell"]: c for c in result["sheets"][0]["cells"]}
        self.assertEqual(cells[self.shift("B6")], {"cell": self.shift("B6"), "kind": "number"})
        self.assertEqual(cells[self.shift("B4")]["text"], "2026-09-30")
        self.assertNotIn("8000000.5", json.dumps(result))

    def test_formula_and_numeric_text_hidden(self):
        wb = load_workbook(self.excel)
        ws = wb.active
        ws["Z1"] = "=SECRET(999999)"
        ws["Z2"] = "1,234.50"
        ws["Z3"] = "(123.50)"
        ws.merge_cells("Z4:AA4")
        ws["Z4"] = "合并表头"
        wb.save(self.excel)
        wb.close()
        result = scan_layout(self.excel)
        content = json.dumps(result, ensure_ascii=False)
        self.assertNotIn("SECRET", content)
        self.assertNotIn("1,234.50", content)
        self.assertNotIn("123.50", content)
        self.assertIn("Z4:AA4", result["sheets"][0]["merged_ranges"])

    def test_bounds_missing_sheet_and_no_silent_truncation(self):
        for area in ("A:A", "A1:XFD200", "A0:B2", "B3:A1"):
            with self.assertRaises(ReportError):
                scan_layout(self.excel, area)
        with self.assertRaises(ReportError):
            scan_layout(self.excel, sheet="missing")
        wb = load_workbook(self.excel)
        wb.active["A1"] = "x" * 401
        wb.save(self.excel)
        wb.close()
        with self.assertRaisesRegex(ReportError, "文字过长"):
            scan_layout(self.excel)

    def test_shifted_layout_end_to_end_and_source_untouched(self):
        proposal_path = self.write_proposal()
        before = self.excel.read_bytes()
        output = self.root / "approved.json"
        result = approve(self.excel, self.config_path, proposal_path, PROJECT / "examples/report_template.docx", output, self.root / "preflight")
        self.assertEqual(result["status"], "approved")
        self.assertFalse((self.root / "preflight/report.docx").exists())
        report = run(self.excel, PROJECT / "examples/report_template.docx", output, self.root / "final")
        self.assertEqual(report["status"], "completed", report["errors"])
        self.assertEqual(len(report["records"]), 32)
        self.assertEqual(len(report["checks"]), 16)
        self.assertEqual(self.excel.read_bytes(), before)

    def test_proposal_preserves_checks_targets_and_limits(self):
        candidate = make_candidate(self.config, self.inventory, self.answer)
        self.assertEqual(candidate["checks"], self.config["checks"])
        self.assertEqual(candidate["format"], self.config["format"])
        for before, after in zip(self.config["fields"], candidate["fields"]):
            self.assertEqual(before.get("target"), after.get("target"))

    def test_period_swap_rejected(self):
        self.reject(lambda a: a["matches"][0].update(header_cell=self.shift("C4"), value_cell=self.shift("C6")))

    def test_amount_hallucination_and_out_of_range_rejected(self):
        self.reject(lambda a: a["matches"][0].update(amount=99))
        self.reject(lambda a: a["matches"][0].update(value_cell="ZZ999"))
        self.reject(lambda a: a["matches"][0].update(label_cell="Z999"))
        self.reject(lambda a: a.update(sheet="invented"))

    def test_duplicates_and_partial_results_rejected(self):
        self.reject(lambda a: a["matches"].append(a["matches"][0]))
        self.reject(lambda a: a["matches"].pop())
        self.reject(lambda a: a.update(unresolved=["cash.end"]))
        self.reject(lambda a: a["matches"][2].update(value_cell=a["matches"][0]["value_cell"]))

    def test_company_unit_metadata_changes_rejected(self):
        self.reject(lambda a: a["metadata"][0].update(cell=self.shift("A6")))
        self.reject(lambda a: a["metadata"].pop())

    def test_label_value_geometry_and_period_group_rejected(self):
        self.reject(lambda a: a["matches"][0].update(value_cell=self.shift("B7")))
        self.reject(lambda a: a["matches"][0].update(value_cell=a["matches"][0]["label_cell"]))
        self.reject(lambda a: a["matches"][1].update(label_cell=self.shift("A7"), value_cell=self.shift("C7")))

    def test_source_changed_or_candidate_tampered_rejected(self):
        proposal = self.write_proposal()
        data = json.loads(proposal.read_text(encoding="utf-8"))
        data["candidate_config"]["checks"][0]["tolerance"] = "1000"
        proposal.write_text(json.dumps(data), encoding="utf-8")
        with self.assertRaisesRegex(ReportError, "被修改"):
            approve(self.excel, self.config_path, proposal, PROJECT / "examples/report_template.docx", self.root / "out.json", self.root / "audit")
        proposal = self.write_proposal()
        with self.excel.open("ab") as stream:
            stream.write(b"changed")
        with self.assertRaisesRegex(ReportError, "已变化"):
            approve(self.excel, self.config_path, proposal, PROJECT / "examples/report_template.docx", self.root / "out.json", self.root / "audit")

    def test_failed_preflight_removes_candidate_keeps_audit(self):
        wb = load_workbook(self.excel)
        wb.active[self.shift("B6")] = 1
        wb.save(self.excel)
        wb.close()
        proposal = self.write_proposal()
        output = self.root / "out.json"
        with self.assertRaisesRegex(ReportError, "预演失败"):
            approve(self.excel, self.config_path, proposal, PROJECT / "examples/report_template.docx", output, self.root / "audit")
        self.assertFalse(output.exists())
        self.assertEqual(json.loads((self.root / "audit/audit.json").read_text(encoding="utf-8"))["status"], "failed")

    def test_no_overwrite_and_accept_required(self):
        path = self.root / "existing.json"
        path.write_text("keep")
        with self.assertRaises(ReportError):
            approve(self.excel, self.config_path, self.write_proposal(), PROJECT / "examples/report_template.docx", path, self.root / "audit")
        self.assertEqual(path.read_text(), "keep")
        result = subprocess.run([sys.executable, "-m", "accounting_report", "approve-mapping"], cwd=PROJECT, capture_output=True)
        self.assertEqual(result.returncode, 2)
        self.assertIn(b"--accept", result.stderr)

    def test_discovery_payload_does_not_send_amounts_and_uses_deepseek(self):
        with patch("accounting_report.layout.request_json", return_value=self.answer) as api:
            result = discover(self.excel, self.config_path)
        args = api.call_args.args
        self.assertEqual(args[1:], (DEEPSEEK_ENDPOINT, DEEPSEEK_MODEL, "DEEPSEEK_API_KEY"))
        self.assertNotIn("8000000.5", args[0][1]["content"])
        self.assertFalse(result["automatic_application"])

    def test_input_change_during_request_rejected(self):
        def changed(*args):
            with self.excel.open("ab") as stream:
                stream.write(b"changed")
            return self.answer
        with patch("accounting_report.layout.request_json", side_effect=changed), self.assertRaisesRegex(ReportError, "发生变化"):
            discover(self.excel, self.config_path)


class DeepSeekTransportTests(unittest.TestCase):
    def call(self, content='{"ok":true}', finish="stop"):
        response = json.dumps({"choices": [{"message": {"content": content}, "finish_reason": finish}]}).encode()
        class Response:
            def __enter__(self): return self
            def __exit__(self, *args): return False
            def read(self, n): return response
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "test-only"}), patch("accounting_report.ai.build_opener") as opener:
            opener.return_value.open.return_value = Response()
            result = request_json([{"role": "user", "content": "JSON"}], DEEPSEEK_ENDPOINT, DEEPSEEK_MODEL, "DEEPSEEK_API_KEY")
            return result, json.loads(opener.return_value.open.call_args.args[0].data)

    def test_deepseek_nonthinking_json_options(self):
        result, payload = self.call()
        self.assertEqual(result, {"ok": True})
        self.assertEqual(payload["thinking"], {"type": "disabled"})
        self.assertEqual(payload["response_format"], {"type": "json_object"})

    def test_truncation_duplicate_keys_and_invalid_json_rejected(self):
        for content, finish in [('{}', 'length'), ('{"ok":1,"ok":2}', 'stop'), ('not json', 'stop'), ('', 'stop')]:
            with self.assertRaises(ReportError):
                self.call(content, finish)

    def test_no_key_no_request(self):
        with patch.dict(os.environ, {}, clear=True), patch("accounting_report.ai.build_opener") as opener:
            with self.assertRaisesRegex(ReportError, "DEEPSEEK_API_KEY"):
                request_json([], DEEPSEEK_ENDPOINT, DEEPSEEK_MODEL, "DEEPSEEK_API_KEY")
        opener.assert_not_called()


if __name__ == "__main__":
    unittest.main()
