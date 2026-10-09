import copy
import csv
import io
import json
import os
import subprocess
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from lxml import etree

from accounting_report.ai import local_suggestions, validate_suggestions, cloud_suggestions
from accounting_report.config import load_config, validate
from accounting_report.errors import ReportError
from accounting_report.extract import ExcelSource
from accounting_report.numbers import amount, display, normalize_label
from accounting_report.pipeline import run, digest, write_audit
from accounting_report.validate import validate_data
from accounting_report.word import WordTemplate, NS, W

PROJECT = Path(__file__).resolve().parents[1]


def minimal_config():
    return {"version": 1, "source": {"sheet": "资产负债表", "unit": "元", "blocks": [
        {"id": "assets", "label_column": "A", "start_row": 3, "end_row": 7, "period_columns": {"end": "B", "begin": "C"}}]},
        "metadata": [{"cell": "A1", "expected": "示例公司"}, {"cell": "B2", "expected": "2026-09-30"}],
        "fields": [{"id": "assets", "source": {"block": "assets", "labels": ["资产总计"], "period": "end"}, "target": {"placeholder": "assets"}},
                   {"id": "liabilities", "source": {"block": "assets", "labels": ["负债合计"], "period": "end"}, "target": {"placeholder": "liabilities"}},
                   {"id": "equity", "source": {"block": "assets", "labels": ["权益合计"], "period": "end"}, "target": {"placeholder": "equity"}}],
        "checks": [{"name": "资产负债平衡", "left": ["assets"], "right": ["liabilities", "equity"], "tolerance": "0.01"}],
        "format": {"unit": "万元", "decimals": 2}}


class NumbersTests(unittest.TestCase):
    def test_numeric_and_grouping(self):
        for value in [1234.5, "1,234.50", "￥1,234.50", "１２３４.５０", Decimal("1234.50")]:
            self.assertEqual(amount(value), Decimal("1234.50"))

    def test_negative(self):
        self.assertEqual(amount("(1,234.50)"), Decimal("-1234.50"))
        self.assertEqual(amount("-12"), Decimal(-12))

    def test_blank_is_not_zero(self):
        for value in [None, "", "  "]:
            with self.assertRaises(ReportError):
                amount(value)
            self.assertEqual(amount(value, True), Decimal(0))
        self.assertEqual(amount(0), Decimal(0))

    def test_invalid_values(self):
        for value in [True, False, float("nan"), float("inf"), "1,23", "1,234,56", "1万元", "10%", "--", "-", "(-1)", "1e3", object()]:
            with self.subTest(value=value), self.assertRaises(ReportError):
                amount(value)

    def test_units_and_round_half_up(self):
        text, converted, rounded = display(Decimal("10050"), "元", "万元", 2)
        self.assertEqual((text, converted, rounded), ("1.01", Decimal("1.005"), Decimal("1.01")))
        self.assertEqual(display(Decimal("-10050"), "元", "万元", 2, True, "parentheses")[0], "(1.01)")
        self.assertEqual(display(Decimal("-0.001"), "元", "元", 2)[0], "0.00")

    def test_label_normalization(self):
        self.assertEqual(normalize_label(" 货 币资金： "), "货币资金")


class FixtureTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.excel = self.root / "source.xlsx"
        self.template = self.root / "template.docx"
        self.config_file = self.root / "mapping.json"
        self.config = minimal_config()
        wb = Workbook()
        ws = wb.active
        ws.title = "资产负债表"
        ws.append(["示例公司"])
        ws.append(["科目", "2026-09-30", "2025-12-31"])
        ws.append(["资产总计", 30000, 25000])
        ws.append(["负债合计", 10000, 5000])
        ws.append(["权益合计", 20000, 20000])
        wb.save(self.excel)
        wb.close()
        doc = Document()
        p = doc.add_paragraph("资产：")
        r = p.add_run("{{ass")
        r.bold = True
        r.font.name = "Arial"
        r = p.add_run("ets}}")
        r.italic = True
        doc.add_paragraph("负债 {{liabilities}}；权益 {{equity}}")
        doc.save(self.template)
        self.save_config()

    def tearDown(self):
        self.tmp.cleanup()

    def save_config(self):
        self.config_file.write_text(json.dumps(self.config, ensure_ascii=False), encoding="utf-8")

    def edit_excel(self, fn):
        wb = load_workbook(self.excel)
        fn(wb.active)
        wb.save(self.excel)
        wb.close()

    def execute(self, dry_run=False):
        self.save_config()
        return run(self.excel, self.template, self.config_file, self.root / "out", dry_run)

    def assert_failure(self, result, substring):
        self.assertEqual(result["status"], "failed")
        self.assertIn(substring, " ".join(result["errors"]))
        self.assertFalse((self.root / "out/report.docx").exists())
        self.assertFalse((self.root / "out/report.pending.docx").exists())
        self.assertTrue((self.root / "out/audit.json").exists())

    def test_end_to_end_format_preservation_and_audit(self):
        before = {str(p): digest(p) for p in [self.excel, self.template, self.config_file]}
        result = self.execute()
        self.assertEqual(result["status"], "completed")
        self.assertEqual(result["checks"][0]["difference"], 0)
        report = self.root / "out/report.docx"
        self.assertEqual(digest(report), result["report_sha256"])
        doc = Document(report)
        self.assertEqual(doc.paragraphs[0].text, "资产：3.00")
        self.assertEqual(doc.paragraphs[1].text, "负债 1.00；权益 2.00")
        self.assertTrue(doc.paragraphs[0].runs[1].bold)
        self.assertTrue(doc.paragraphs[0].runs[2].italic)
        self.assertEqual(doc.paragraphs[0].runs[1].font.name, "Arial")
        for p in [self.excel, self.template, self.config_file]:
            self.assertEqual(digest(p), before[str(p)])
        with ZipFile(self.template) as a, ZipFile(report) as b:
            self.assertEqual(set(a.namelist()), set(b.namelist()))
            for name in a.namelist():
                if name != "word/document.xml":
                    self.assertEqual(a.read(name), b.read(name), name)
            ra = etree.fromstring(a.read("word/document.xml"))
            rb = etree.fromstring(b.read("word/document.xml"))
            self.assertEqual([etree.tostring(x) for x in ra.xpath("//w:rPr", namespaces=NS)],
                             [etree.tostring(x) for x in rb.xpath("//w:rPr", namespaces=NS)])
        logged = json.loads((self.root / "out/audit.json").read_text(encoding="utf-8"))
        self.assertEqual(logged["records"][0]["value"], "30000")
        self.assertEqual(logged["records"][0]["cell"], "B3")
        with (self.root / "out/audit.csv").open(encoding="utf-8-sig", newline="") as f:
            self.assertEqual(len(list(csv.DictReader(f))), 3)

    def test_dry_run_validates_without_report(self):
        result = self.execute(True)
        self.assertEqual(result["status"], "validated")
        self.assertEqual(len(result["word_changes"]), 3)
        self.assertFalse((self.root / "out/report.docx").exists())

    def test_balance_failure(self):
        self.edit_excel(lambda s: setattr(s["B3"], "value", 30001))
        self.assert_failure(self.execute(), "差额 1")

    def test_missing_blank_and_missing_label(self):
        self.edit_excel(lambda s: setattr(s["B4"], "value", None))
        self.assert_failure(self.execute(), "金额为空")

    def test_missing_label(self):
        self.edit_excel(lambda s: setattr(s["A4"], "value", "其他负债"))
        self.assert_failure(self.execute(), "匹配 0 行")

    def test_duplicate_label(self):
        self.edit_excel(lambda s: s.append(["资产总计", 30000, 25000]))
        self.assert_failure(self.execute(), "匹配 2 行")

    def test_alias_matches_multiple_rows_are_ambiguous(self):
        self.config["fields"][0]["source"]["labels"].append("资产合计")
        self.edit_excel(lambda s: s.append(["资产合计", 30000, 25000]))
        self.assert_failure(self.execute(), "匹配 2 行")

    def test_wrong_company_or_period(self):
        self.edit_excel(lambda s: setattr(s["B2"], "value", "2026-08-31"))
        self.assert_failure(self.execute(), "来源标识不符")

    def test_missing_sheet(self):
        self.config["source"]["sheet"] = "不存在"
        self.assert_failure(self.execute(), "找不到工作表")

    def test_formula_without_cache_is_blocked_even_if_blank_zero(self):
        self.edit_excel(lambda s: setattr(s["B3"], "value", "=B4+B5"))
        self.config["fields"][0]["blank_zero"] = True
        self.assert_failure(self.execute(), "没有缓存结果")

    def test_excel_error(self):
        self.edit_excel(lambda s: setattr(s["B3"], "value", "#DIV/0!"))
        self.assert_failure(self.execute(), "Excel 错误值")

    def test_bounds(self):
        self.config["fields"][0]["max"] = "20000"
        self.assert_failure(self.execute(), "不符合 max")

    def test_valid_zero(self):
        self.edit_excel(lambda s: (setattr(s["B4"], "value", 0), setattr(s["B3"], "value", 20000)))
        self.assertEqual(self.execute()["records"][1]["formatted"], "0.00")

    def test_explicit_blank_zero_warning(self):
        self.edit_excel(lambda s: (setattr(s["B4"], "value", None), setattr(s["B3"], "value", 20000)))
        self.config["fields"][1]["blank_zero"] = True
        result = self.execute()
        self.assertEqual(result["status"], "completed")
        self.assertIn("blank_zero", " ".join(result["warnings"]))

    def test_coordinate_mapping_and_drift_guard(self):
        self.config["fields"][0]["source"] = {"cell": "B3", "label_cell": "A3", "labels": ["资产总计"]}
        self.assertEqual(self.execute()["status"], "completed")

    def test_coordinate_drift(self):
        self.config["fields"][0]["source"] = {"cell": "B4", "label_cell": "A4", "labels": ["资产总计"]}
        self.assert_failure(self.execute(), "坐标科目校验失败")

    def test_word_unmapped_placeholder(self):
        d = Document(self.template)
        d.add_paragraph("{{unknown}}")
        d.save(self.template)
        self.assert_failure(self.execute(), "未映射占位符")

    def test_word_missing_or_duplicate_target(self):
        self.config["fields"][0]["target"]["occurrences"] = 2
        self.assert_failure(self.execute(), "配置要求 2 次")

    def test_malformed_placeholder(self):
        d = Document(self.template)
        d.add_paragraph("{{中文变量}}")
        d.save(self.template)
        self.assert_failure(self.execute(), "格式错误的占位符")

    def test_header_footer_and_nested_table(self):
        d = Document(self.template)
        d.sections[0].header.paragraphs[0].text = "资产 {{assets}}"
        d.sections[0].footer.paragraphs[0].text = "负债 {{liabilities}}"
        table = d.add_table(rows=1, cols=1)
        nested = table.cell(0, 0).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "{{equity}}"
        d.save(self.template)
        for f in self.config["fields"]:
            f["target"]["occurrences"] = 2
        result = self.execute()
        self.assertEqual(result["status"], "completed")
        self.assertEqual(len(result["word_changes"]), 6)
        doc = Document(self.root / "out/report.docx")
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, "资产 3.00")
        self.assertEqual(doc.sections[0].footer.paragraphs[0].text, "负债 1.00")

    def test_cell_target_empty_preserves_style(self):
        d = Document(self.template)
        d.paragraphs[0].clear()
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "资产总计"
        r = t.cell(0, 1).paragraphs[0].add_run("")
        r.bold = True
        d.save(self.template)
        self.config["fields"][0]["target"] = {"table": 0, "row": 0, "cell": 1, "expected_text": ""}
        result = self.execute()
        self.assertEqual(result["status"], "completed")
        cell = Document(self.root / "out/report.docx").tables[0].cell(0, 1)
        self.assertEqual(cell.text, "3.00")
        self.assertTrue(cell.paragraphs[0].runs[0].bold)

    def test_cell_expected_text_prevents_overwrite(self):
        d = Document(self.template)
        d.paragraphs[0].clear()
        d.add_table(rows=1, cols=1).cell(0, 0).text = "已填金额"
        d.save(self.template)
        self.config["fields"][0]["target"] = {"table": 0, "row": 0, "cell": 0, "expected_text": ""}
        self.assert_failure(self.execute(), "原文不符")

    def test_tracked_changes_blocked(self):
        d = Document(self.template)
        ins = OxmlElement("w:ins")
        d.paragraphs[0]._p.append(ins)
        d.save(self.template)
        self.assert_failure(self.execute(), "含修订")

    def test_field_display_placeholder_blocked(self):
        d = Document(self.template)
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), "begin")
        d.paragraphs[0].runs[0]._r.append(field)
        d.save(self.template)
        self.assert_failure(self.execute(), "含 Word 域")

    def test_placeholder_cannot_cross_tab(self):
        d = Document(self.template)
        d.paragraphs[0].clear()
        d.paragraphs[0].add_run("{{ass\tets}}")
        d.save(self.template)
        self.assert_failure(self.execute(), "出现 0 次")

    def test_duplicate_word_target(self):
        self.config["fields"][1]["target"] = {"placeholder": "assets"}
        self.assert_failure(self.execute(), "Word 目标重复")

    def test_period_selection(self):
        for f in self.config["fields"]:
            f["source"]["period"] = "begin"
        result = self.execute()
        self.assertEqual(result["status"], "completed")
        self.assertEqual(result["records"][0]["cell"], "C3")
        self.assertEqual(result["records"][0]["formatted"], "2.50")

    def test_tolerance_boundary_passes(self):
        self.edit_excel(lambda s: setattr(s["B3"], "value", "30000.01"))
        self.assertEqual(self.execute()["status"], "completed")

    def test_document_protection_blocked(self):
        d = Document(self.template)
        protection = OxmlElement("w:documentProtection")
        protection.set(qn("w:enforcement"), "1")
        d.settings.element.append(protection)
        d.save(self.template)
        self.assert_failure(self.execute(), "编辑保护")

    def test_existing_output_never_overwritten(self):
        out = self.root / "out"
        out.mkdir()
        sentinel = out / "keep.txt"
        sentinel.write_text("unchanged")
        with self.assertRaises(ReportError):
            self.execute()
        self.assertEqual(sentinel.read_text(), "unchanged")

    def test_report_publication_failure_is_logged(self):
        with patch("accounting_report.pipeline.os.replace", side_effect=PermissionError("blocked")):
            self.assert_failure(self.execute(), "报告发布失败")
        logged = json.loads((self.root / "out/audit.json").read_text(encoding="utf-8"))
        self.assertEqual(logged["status"], "failed")

    def test_source_modified_during_run_blocks_publication(self):
        original = WordTemplate.fill
        def modified_fill(word, config, records):
            logs = original(word, config, records)
            self.config_file.write_text(self.config_file.read_text(encoding="utf-8") + " ", encoding="utf-8")
            return logs
        with patch.object(WordTemplate, "fill", modified_fill):
            self.assert_failure(self.execute(), "输入文件发生变化")

    def test_whitespace_blank_zero_is_audited(self):
        self.edit_excel(lambda s: (setattr(s["B4"], "value", " "), setattr(s["B3"], "value", 20000)))
        self.config["fields"][1]["blank_zero"] = True
        result = self.execute()
        self.assertTrue(result["records"][1]["blank_as_zero"])

    def test_csv_formula_text_is_neutralized(self):
        out = self.root / "csv-test"
        out.mkdir()
        write_audit(out, {"records": [{"id": "test", "formula": "=SUM(B3:B5)", "label": "@unsafe"}]})
        with (out / "audit.csv").open(encoding="utf-8-sig", newline="") as f:
            row = next(csv.DictReader(f))
        self.assertEqual(row["formula"], "'=SUM(B3:B5)")
        self.assertEqual(row["label"], "'@unsafe")

    def test_numeric_merged_child_is_not_silently_zero(self):
        self.edit_excel(lambda s: s.merge_cells("B3:C3"))
        self.config["fields"][0]["source"] = {"cell": "C3", "label_cell": "A3", "labels": ["资产总计"]}
        self.assert_failure(self.execute(), "金额为空")

    def test_invalid_config_still_logs(self):
        self.config["format"]["decimals"] = -1
        self.assert_failure(self.execute(), "decimals")

    def test_missing_input_still_logs(self):
        self.excel.unlink()
        self.assert_failure(self.execute(), "FileNotFoundError")

    def test_cli_success_and_failure_codes(self):
        env = {**os.environ, "PYTHONUTF8": "1"}
        args = [sys.executable, "-m", "accounting_report", "run", "--excel", str(self.excel),
                "--template", str(self.template), "--config", str(self.config_file), "--output-dir", str(self.root / "cli-out")]
        result = subprocess.run(args, cwd=PROJECT, env=env, capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0, result.stderr)
        result = subprocess.run(args, cwd=PROJECT, env=env, capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 2)


class ConfigurationTests(unittest.TestCase):
    def test_unknown_keys_and_empty_checks_rejected(self):
        for change in [lambda c: c.update({"typo": True}), lambda c: c.update({"checks": []}),
                       lambda c: c["fields"].append(copy.deepcopy(c["fields"][0])),
                       lambda c: c["checks"][0].update({"right": ["missing"]}),
                       lambda c: c["format"].update({"decimals": True}),
                       lambda c: c["fields"][0].update({"blank_zero": "false"})]:
            c = minimal_config()
            change(c)
            with self.assertRaises(ReportError):
                validate(c)

    def test_duplicate_json_keys(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "bad.json"
            p.write_text('{"version":1,"version":1}')
            with self.assertRaisesRegex(ReportError, "键重复"):
                load_config(p)

    def test_unavailable_check_is_not_pass(self):
        checks, errors = validate_data(minimal_config(), [])
        self.assertEqual(checks[0]["status"], "unavailable")
        self.assertTrue(errors)


class AITests(unittest.TestCase):
    def setUp(self):
        self.c = minimal_config()
        self.labels = [{"block": "assets", "cell": "A3", "label": "资产总计"}]

    def test_local_no_network(self):
        with patch("accounting_report.ai.build_opener", side_effect=AssertionError("network prohibited")):
            result = local_suggestions(self.c, self.labels)
        self.assertFalse(result["automatic_application"])
        self.assertEqual(result["suggestions"][0]["candidates"][0]["similarity"], 1)

    def test_ai_whitelist_valid_and_hallucination_rejected(self):
        item = {"field": "assets", "block": "assets", "cell": "A3", "reason": "名称一致"}
        self.assertFalse(validate_suggestions({"suggestions": [item]}, self.c, self.labels)["automatic_application"])
        for bad in [{**item, "cell": "A999"}, {**item, "field": "invented"}, {**item, "block": "other"}, {**item, "amount": 100}]:
            with self.assertRaises(ReportError):
                validate_suggestions({"suggestions": [bad]}, self.c, self.labels)

    def test_ai_requires_key_and_https(self):
        with self.assertRaises(ReportError):
            cloud_suggestions(self.c, self.labels, "http://example.com", "model")
        with patch.dict(os.environ, {}, clear=True), self.assertRaises(ReportError):
            cloud_suggestions(self.c, self.labels, "https://example.com/chat/completions", "model")

    def test_ai_request_sends_labels_only(self):
        item = {"field": "assets", "block": "assets", "cell": "A3", "reason": "名称一致"}
        response = json.dumps({"choices": [{"message": {"content": json.dumps({"suggestions": [item]})}}]}).encode()
        class Response:
            def __enter__(self): return self
            def __exit__(self, *args): return False
            def read(self, n): return response
        with patch.dict(os.environ, {"REPORT_AI_API_KEY": "test-key"}), patch("accounting_report.ai.build_opener") as opener:
            opener.return_value.open.return_value = Response()
            result = cloud_suggestions(self.c, self.labels, "https://example.com/v1/chat/completions", "test-model")
            request = opener.return_value.open.call_args.args[0]
            data = json.loads(request.data)
            sent = json.loads(data["messages"][1]["content"])
            self.assertEqual(set(sent), {"fields", "candidate_labels"})
            self.assertNotIn("示例公司", request.data.decode())
            self.assertFalse(result["automatic_application"])


class ShippedExampleTests(unittest.TestCase):
    def test_bundled_example_complete(self):
        with tempfile.TemporaryDirectory() as td:
            result = run(PROJECT / "examples/balance_sheet.xlsx", PROJECT / "examples/report_template.docx",
                         PROJECT / "examples/mapping.json", Path(td) / "output")
            self.assertEqual(result["status"], "completed", result["errors"])
            self.assertEqual(len(result["records"]), 32)
            self.assertEqual(len(result["checks"]), 16)
            self.assertTrue(all(x["status"] == "passed" for x in result["checks"]))
            self.assertTrue(all(x["difference"] == 0 for x in result["checks"]))
            values = {x["id"]: x for x in result["records"]}
            self.assertEqual(values["assets.end"]["value"], Decimal("8000000.5"))
            self.assertEqual(values["assets.end"]["formatted"], "800.00")
            self.assertEqual(values["assets.begin"]["formatted"], "680.00")
            self.assertFalse(WordTemplate(Path(td) / "output/report.docx").inventory()["placeholders"])


if __name__ == "__main__":
    unittest.main()
