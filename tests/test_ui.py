import base64
import io
import json
import os
import tempfile
import threading
import unittest
from http.client import HTTPConnection
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook

from accounting_report.config import load_config
from accounting_report.errors import ReportError
from accounting_report.layout import discover
from accounting_report.ui import create_server, App, unpack_file, MAX_FILE, MAX_BODY

PROJECT = Path(__file__).resolve().parents[1]


def encoded(path):
    return {"name": path.name, "data": base64.b64encode(path.read_bytes()).decode()}


def mock_discover(excel, config_path):
    config = load_config(config_path)
    wb = load_workbook(excel, data_only=False)
    ws = wb[config["source"]["sheet"]]
    answer = {"sheet": ws.title, "metadata": [{"original_cell": m["cell"], "cell": m["cell"]} for m in config["metadata"]],
              "matches": [], "unresolved": []}
    for field in config["fields"]:
        src = field["source"]
        b = next(b for b in config["source"]["blocks"] if b["id"] == src["block"])
        row = next(r for r in range(b["start_row"], b["end_row"] + 1) if ws[f"{b['label_column']}{r}"].value in src["labels"])
        col = b["period_columns"][src["period"]]
        answer["matches"].append({"field": field["id"], "label_cell": f"{b['label_column']}{row}",
             "value_cell": f"{col}{row}", "header_cell": f"{col}4", "reason": "科目和期间表头对应"})
    wb.close()
    with patch("accounting_report.layout.request_json", return_value=answer):
        return discover(excel, config_path)


class UITests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.server = create_server(PROJECT, self.root, 0)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()
        self.temp.cleanup()

    def request(self, path, data=None, headers=None):
        conn = HTTPConnection("127.0.0.1", self.server.server_port, timeout=15)
        defaults = {"X-App-Token": self.server.app.token, "Content-Type": "application/json"}
        defaults.update(headers or {})
        conn.request("POST" if data is not None else "GET", path, body=json.dumps(data) if data is not None else None, headers=defaults)
        response = conn.getresponse()
        body = response.read()
        result = (response.status, body, dict(response.getheaders()))
        conn.close()
        return result

    def post(self, path, data, headers=None):
        status, body, _ = self.request(path, data, headers)
        return status, json.loads(body)

    def test_page_assets_and_catalogue(self):
        status, body, headers = self.request("/")
        self.assertEqual(status, 200)
        self.assertIn("上传表格".encode(), body)
        self.assertIn(self.server.app.token.encode(), body)
        self.assertEqual(headers["Cache-Control"], "no-store")
        self.assertIn("frame-ancestors 'none'", headers["Content-Security-Policy"])
        for asset in ("/app.js", "/style.css"):
            self.assertEqual(self.request(asset)[0], 200)
        status, body, _ = self.request("/api/templates")
        catalogue = json.loads(body)
        self.assertEqual(catalogue["templates"][0]["id"], "example")
        self.assertNotIn("word", catalogue["templates"][0])

    def test_uploaded_excel_generates_downloadable_report_and_logs(self):
        status, data = self.post("/api/run", {"template": "example", "excel": encoded(PROJECT / "examples/balance_sheet.xlsx")})
        self.assertEqual(status, 200)
        self.assertEqual(data["status"], "completed")
        self.assertEqual((data["records"], data["checks"]), (32, 16))
        prefix = f"/download/{data['id']}/"
        status, body, _ = self.request(prefix + "report-and-audit.zip?token=" + self.server.app.token)
        self.assertEqual(status, 200)
        with ZipFile(io.BytesIO(body)) as archive:
            self.assertEqual(set(archive.namelist()), {"report.docx", "audit.json", "audit.csv"})
            self.assertEqual(json.loads(archive.read("audit.json"))["status"], "completed")

    def test_demo_no_key_no_cloud(self):
        with patch.dict(os.environ, {}, clear=True), patch("accounting_report.ui.discover", side_effect=AssertionError("cloud")):
            status, data = self.post("/api/run", {"template": "example", "demo": True})
        self.assertEqual((status, data["status"]), (200, "completed"))

    def test_failed_validation_blocks_report_but_logs_download(self):
        excel = self.root / "wrong.xlsx"
        wb = load_workbook(PROJECT / "examples/balance_sheet.xlsx")
        wb["资产负债表"]["B3"] = "另一家公司"
        wb.save(excel)
        wb.close()
        status, data = self.post("/api/run", {"template": "example", "excel": encoded(excel)})
        self.assertEqual((status, data["status"]), (200, "failed"))
        self.assertTrue(data["errors"])
        prefix = f"/download/{data['id']}/"
        self.assertEqual(self.request(prefix + "report.docx?token=" + self.server.app.token)[0], 404)
        self.assertEqual(self.request(prefix + "audit.json?token=" + self.server.app.token)[0], 200)

    def test_custom_template_saved_reloaded_and_used(self):
        status, data = self.post("/api/template", {"name": "我的报告", "word": encoded(PROJECT / "examples/report_template.docx"),
                                                "config": encoded(PROJECT / "examples/mapping.json")})
        self.assertEqual(status, 200)
        restored = App(PROJECT, self.root)
        self.assertEqual(restored.templates[data["id"]]["name"], "我的报告")
        status, result = self.post("/api/run", {"template": data["id"], "demo": True})
        self.assertEqual(result["status"], "completed")

    def test_invalid_template_not_registered(self):
        status, _ = self.post("/api/template", {"name": "无效模板", "word": {"name": "x.docx", "data": base64.b64encode(b"bad").decode()},
                                                "config": encoded(PROJECT / "examples/mapping.json")})
        self.assertEqual(status, 400)
        self.assertEqual(len(self.server.app.templates), 1)

    def test_token_origin_and_host_validation(self):
        payload = {"template": "example", "demo": True}
        for headers in ({"X-App-Token": "wrong"}, {"Origin": "https://evil.example"}, {"Host": "evil.example"}):
            status, _ = self.post("/api/run", payload, headers)
            self.assertEqual(status, 403)
        self.assertEqual(self.request("/", headers={"Host": "evil.example"})[0], 403)

    def test_no_arbitrary_paths_or_download_without_token(self):
        self.assertEqual(self.request("/../README.md")[0], 404)
        self.assertEqual(self.request("/download/nope/report.docx")[0], 403)
        self.assertEqual(self.request("/download/nope/mapping.json?token=" + self.server.app.token)[0], 404)
        status, _ = self.post("/api/run", {"template": "../examples", "demo": True})
        self.assertEqual(status, 400)

    def test_no_key_ai_fails_before_network(self):
        with patch.dict(os.environ, {}, clear=True), patch("accounting_report.ui.discover") as api:
            status, data = self.post("/api/run", {"template": "example", "demo": True, "use_ai": True})
        self.assertEqual(status, 400)
        self.assertIn("DEEPSEEK_API_KEY", data["error"])
        api.assert_not_called()

    def test_ai_review_then_confirmation_generates_report(self):
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "test"}), patch("accounting_report.ui.discover", side_effect=mock_discover):
            status, data = self.post("/api/run", {"template": "example", "demo": True, "use_ai": True})
        self.assertEqual((status, data["status"]), (200, "needs_review"))
        self.assertEqual(len(data["matches"]), 32)
        self.assertIn("label", data["matches"][0])
        self.assertEqual(self.request(f"/download/{data['id']}/report.docx?token=" + self.server.app.token)[0], 404)
        self.assertEqual(self.post("/api/approve", {"id": data["id"], "accept": False})[0], 400)
        status, result = self.post("/api/approve", {"id": data["id"], "accept": True})
        self.assertEqual((status, result["status"]), (200, "completed"))
        self.assertEqual(self.post("/api/approve", {"id": data["id"], "accept": True})[0], 400)

    def test_ai_confirmation_rejects_modified_proposal(self):
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "test"}), patch("accounting_report.ui.discover", side_effect=mock_discover):
            _, data = self.post("/api/run", {"template": "example", "demo": True, "use_ai": True})
        # Alter the proposal, not the trusted generated mapping.
        directory = self.server.app.jobs[data["id"]]["directory"]
        proposal = directory / "proposal.json"
        content = json.loads(proposal.read_text(encoding="utf-8"))
        content["candidate_config"]["checks"][0]["tolerance"] = "100"
        proposal.write_text(json.dumps(content), encoding="utf-8")
        status, result = self.post("/api/approve", {"id": data["id"], "accept": True})
        self.assertEqual((status, result["status"]), (200, "failed"))
        self.assertTrue(result["errors"])

    def test_upload_constraints(self):
        for item in ({"name": "a.xls", "data": "YQ=="}, {"name": "a.xlsx", "data": "??"}, {"name": "a.xlsx", "data": ""}):
            with self.assertRaises(ReportError):
                unpack_file(item, (".xlsx",))
        for payload in ({"template": "example", "demo": "yes"}, {"template": "missing", "demo": True}):
            self.assertEqual(self.post("/api/run", payload)[0], 400)
        with self.assertRaises(ReportError):
            unpack_file({"name": "large.xlsx", "data": "A" * (MAX_FILE * 4 // 3 + 8)}, (".xlsx",))
        self.assertEqual(self.post("/api/run", {"template": "example", "demo": True},
                                   {"Content-Length": str(MAX_BODY + 1)})[0], 400)


if __name__ == "__main__":
    unittest.main()
