from openpyxl import load_workbook

from .errors import ReportError
from .numbers import amount, normalize_label


class ExcelSource:
    def __init__(self, path, config):
        self.config = config
        self.formulas = load_workbook(path, data_only=False, keep_links=False)
        try:
            self.cached = load_workbook(path, data_only=True, keep_links=False)
            sheet = config["source"]["sheet"]
            if sheet not in self.formulas.sheetnames:
                raise ReportError(f"找不到工作表：{sheet}")
            self.ws = self.formulas[sheet]
            self.values = self.cached[sheet]
            self.blocks = {b["id"]: b for b in config["source"]["blocks"]}
            for m in config["metadata"]:
                actual = self.ws[m["cell"]].value
                if str(actual).strip() != m["expected"].strip():
                    raise ReportError(f"来源标识不符 {sheet}!{m['cell']}：期望 {m['expected']!r}，实际 {actual!r}")
        except Exception:
            self.close()
            raise

    def close(self):
        self.formulas.close()
        if hasattr(self, "cached"):
            self.cached.close()

    def labels(self):
        result = []
        for b in self.blocks.values():
            for row in range(b["start_row"], b["end_row"] + 1):
                cell = self.ws[f"{b['label_column']}{row}"]
                if cell.value is not None:
                    result.append({"block": b["id"], "cell": cell.coordinate, "label": str(cell.value)})
        return result

    def read(self, field):
        src = field["source"]
        labels = {normalize_label(x) for x in src["labels"]}
        if "cell" in src:
            label_cell = src["label_cell"]
            coord = src["cell"]
            if normalize_label(self.ws[label_cell].value) not in labels:
                raise ReportError(f"坐标科目校验失败：{label_cell} 为 {self.ws[label_cell].value!r}")
        else:
            b = self.blocks[src["block"]]
            matches = [r for r in range(b["start_row"], b["end_row"] + 1)
                       if self.ws[f"{b['label_column']}{r}"].value is not None
                       and normalize_label(self.ws[f"{b['label_column']}{r}"].value) in labels]
            if len(matches) != 1:
                raise ReportError(f"科目 {src['labels']} 在 {b['id']} 匹配 {len(matches)} 行，要求唯一；行号 {matches}")
            label_cell = f"{b['label_column']}{matches[0]}"
            coord = f"{b['period_columns'][src['period']]}{matches[0]}"
        c = self.ws[coord]
        formula = c.value if c.data_type == "f" else None
        raw = self.values[coord].value if formula else c.value
        if c.data_type == "e" or self.values[coord].data_type == "e":
            raise ReportError(f"Excel 错误值：{coord} {raw}")
        if formula and raw is None:
            raise ReportError(f"公式 {coord} 没有缓存结果；请在 Excel 中重算并保存后重试")
        if formula and "[" in formula:
            raise ReportError(f"{coord} 含外部工作簿引用，须转成经核实的数值")
        value = amount(raw, field.get("blank_zero", False))
        return {"id": field["id"], "label": str(self.ws[label_cell].value),
                "label_cell": label_cell, "sheet": self.ws.title, "cell": coord,
                "period": src.get("period", "coordinate"), "raw": raw,
                "formula": formula, "value": value,
                "blank_as_zero": raw is None or (isinstance(raw, str) and not raw.strip())}
