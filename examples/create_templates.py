"""Re-create the replaceable sample Word template and mapping (not a real report).

Run from project root: python examples/create_templates.py
Requires python-docx; the Excel sample is already included.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
ITEMS = [
    ("cash", "货币资金", "assets", 6),
    ("receivables", "应收账款", "assets", 7),
    ("inventory", "存货", "assets", 8),
    ("current_assets", "流动资产合计", "assets", 9),
    ("fixed_assets", "固定资产", "assets", 10),
    ("noncurrent_assets", "非流动资产合计", "assets", 11),
    ("assets", "资产总计", "assets", 12),
    ("short_loan", "短期借款", "liabilities", 6),
    ("payables", "应付账款", "liabilities", 7),
    ("current_liabilities", "流动负债合计", "liabilities", 8),
    ("long_loan", "长期借款", "liabilities", 9),
    ("liabilities", "负债合计", "liabilities", 10),
    ("capital", "实收资本", "liabilities", 11),
    ("retained", "未分配利润", "liabilities", 12),
    ("equity", "所有者权益合计", "liabilities", 13),
    ("liabilities_equity", "负债和所有者权益总计", "liabilities", 14),
]


def make_config():
    fields = []
    for ident, label, block, row in ITEMS:
        for period in ["end", "begin"]:
            identifier = f"{ident}.{period}"
            target = {"placeholder": identifier}
            if period == "end" and ident in ["assets", "liabilities", "equity"]:
                target["occurrences"] = 2
            fields.append({"id": identifier, "source": {"block": block, "labels": [label], "period": period}, "target": target})
    checks = []
    equations = [
        ("流动资产分项", ["current_assets"], ["cash", "receivables", "inventory"]),
        ("非流动资产分项", ["noncurrent_assets"], ["fixed_assets"]),
        ("资产合计", ["assets"], ["current_assets", "noncurrent_assets"]),
        ("流动负债分项", ["current_liabilities"], ["short_loan", "payables"]),
        ("负债合计", ["liabilities"], ["current_liabilities", "long_loan"]),
        ("权益合计", ["equity"], ["capital", "retained"]),
        ("负债和权益合计", ["liabilities_equity"], ["liabilities", "equity"]),
        ("资产负债平衡", ["assets"], ["liabilities", "equity"]),
    ]
    for period, title in [("end", "期末"), ("begin", "期初")]:
        for name, left, right in equations:
            checks.append({"name": title + name, "left": [f"{x}.{period}" for x in left],
                           "right": [f"{x}.{period}" for x in right], "tolerance": "0.01"})
    return {"version": 1,
            "source": {"sheet": "资产负债表", "unit": "元", "blocks": [
                {"id": "assets", "label_column": "A", "start_row": 6, "end_row": 12, "period_columns": {"end": "B", "begin": "C"}},
                {"id": "liabilities", "label_column": "E", "start_row": 6, "end_row": 14, "period_columns": {"end": "F", "begin": "G"}}]},
            "metadata": [{"cell": "B3", "expected": "示例有限公司"}, {"cell": "F3", "expected": "元"},
                         {"cell": "B4", "expected": "2026-09-30"}, {"cell": "C4", "expected": "2025-12-31"},
                         {"cell": "F4", "expected": "2026-09-30"}, {"cell": "G4", "expected": "2025-12-31"}],
            "fields": fields, "checks": checks,
            "format": {"unit": "万元", "decimals": 2, "grouping": True, "negative": "parentheses"}}


def style_run(run, bold=False):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(10)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def create_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.7), Cm(1.7)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    for name in ["Normal", "Title", "Heading 1"]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    # The python-docx base template may carry a blue border in Title/Subtitle.
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph("示例有限公司会计报告", style="Title")
    p.paragraph_format.space_after = Pt(8)
    p = doc.add_paragraph("报告日 2026年9月30日    对比日 2025年12月31日    金额单位 万元")
    for r in p.runs:
        style_run(r)
    doc.add_paragraph("本报告列示公司期末和期初的资产、负债及所有者权益余额。主要余额如下，详细金额列于下表。")
    p = doc.add_paragraph()
    # Intentionally split a placeholder across runs to exercise format preservation.
    for text, bold in [("截至报告日，资产总额为 ", False), ("{{assets.", True), ("end}}", True),
                       (" 万元，负债总额为 {{liabilities.end}} 万元，所有者权益为 {{equity.end}} 万元。", False)]:
        style_run(p.add_run(text), bold)
    p = doc.add_paragraph("资产负债余额明细", style="Heading 1")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Cm(6.5), Cm(5.25), Cm(5.25)]
    for col, width in zip(table.columns, widths):
        col.width = width
    for cell, value in zip(table.rows[0].cells, ["科目", "期末余额", "期初余额"]):
        cell.text = value
    for ident, label, block, row in ITEMS:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = "{{" + ident + ".end}}"
        cells[2].text = "{{" + ident + ".begin}}"
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = widths[ci]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ["top", "left", "bottom", "right"]:
                e = OxmlElement("w:" + edge)
                for k, v in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                    e.set(qn("w:" + k), v)
                borders.append(e)
            tcpr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ["top", "bottom", "left", "right"]:
                e = OxmlElement("w:" + edge)
                e.set(qn("w:w"), "80")
                e.set(qn("w:type"), "dxa")
                margins.append(e)
            tcpr.append(margins)
            if ri == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E7EDF4")
                tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else (WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.RIGHT)
                for r in p.runs:
                    style_run(r, ri == 0 or "合计" in row.cells[0].text or "总计" in row.cells[0].text)
    doc.add_paragraph("金额按万元保留两位小数列示，分项与合计可能出现显示舍入差额。")
    doc.save(ROOT / "report_template.docx")


if __name__ == "__main__":
    (ROOT / "mapping.json").write_text(json.dumps(make_config(), ensure_ascii=False, indent=2), encoding="utf-8")
    create_docx()
    print("已生成 report_template.docx 和 mapping.json；示例 Excel 已随项目提供")
